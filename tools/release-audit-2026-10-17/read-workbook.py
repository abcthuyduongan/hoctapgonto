"""Read-only extraction; never saves or modifies the input workbook."""
import json
import pathlib
import sys
import openpyxl
from openpyxl.utils import get_column_letter

source = pathlib.Path(sys.argv[1])
target = pathlib.Path(sys.argv[2])
workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
result = []
for sheet in workbook:
    rows = []
    for index, row in enumerate(sheet.iter_rows(), 1):
        cells = [{"col": get_column_letter(col), "value": str(cell.value)}
                 for col, cell in enumerate(row, 1) if cell.value is not None]
        if cells:
            rows.append({"row": index, "cells": cells})
    result.append({"sheet": sheet.title, "rows": rows})
workbook.close()
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Read {len(result)} sheets; source workbook was not saved.")
